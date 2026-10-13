from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = Path(__file__).resolve().parent.parent / "docs" / "Week_1_and_2_Graduation_Project_Report.docx"

REPO = "https://github.com/munthermohammed064-pixel/skillstream-learning-platform"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)

doc = Document()

for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.5)
    s.left_margin = s.right_margin = Cm(2.5)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size in (("Heading 1", 15), ("Heading 2", 13)):
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
    st.element.rPr.rFonts.set(qn("w:hAnsiTheme"), "")
    st.element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    st.element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = ACCENT
    st.paragraph_format.space_before = Pt(14)
    st.paragraph_format.space_after = Pt(6)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def para(text, bold=False, italic=False, size=None, align=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def body(text):
    p = para(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F3A5F")
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
            if n % 2:
                shade(cells[i], "EEF2F7")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
            for p in row.cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(11)
    doc.add_paragraph()


# ---------- Title page ----------
for _ in range(3):
    doc.add_paragraph()
para("ALTINBAŞ UNIVERSITY", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Faculty of Engineering and Natural Sciences", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Department of Software Engineering", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
para("Graduation Project", italic=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
title = para("SkillStream: Multi-Role Online Learning Platform", bold=True, size=22,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
title.runs[0].font.color.rgb = ACCENT
para("Progress Report: Week 1 and Week 2", size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=60)

info = doc.add_table(rows=0, cols=2)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, v in (
    ("Course", "Graduation Project"),
    ("Institution", "Altınbaş University"),
    ("Instructor", "Dr. F. Kuzey Edes Huyal"),
    ("Prepared by", "MUNTHER ISKENDER  (210513647)"),
    ("Reporting Period", "Week 1 – Week 2"),
    ("Date", "October 2026"),
):
    cells = info.add_row().cells
    cells[0].text, cells[1].text = "", ""
    cells[0].paragraphs[0].add_run(k).bold = True
    cells[1].paragraphs[0].add_run(v)
    cells[0].width, cells[1].width = Cm(4.5), Cm(8.5)

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- Week 1 ----------
doc.add_heading("Part I: Week 1 – Project Definition", level=1)

doc.add_heading("1.1 Project Title", level=2)
body("SkillStream: Multi-Role Online Learning Platform.")

doc.add_heading("1.2 Problem Statement", level=2)
body(
    "Online learning platforms frequently encounter difficulties in organising and distributing "
    "permissions among administrators, instructors, and learners. When responsibilities are not "
    "clearly separated, instructors depend on administrative staff for routine content updates, "
    "administrators lose visibility over the quality of published material, and students are "
    "presented with fragmented interfaces that hinder course discovery and consumption."
)
body(
    "Instructors therefore require a straightforward means of publishing and maintaining their "
    "own course content, whereas students require an intuitive environment in which they can "
    "browse the course catalogue, enrol in courses, and watch lessons without interruption. "
    "SkillStream addresses this problem by providing a single web-based platform built around "
    "a clearly defined role-based access model comprising three roles: Admin, Teacher, and Student."
)

doc.add_heading("1.3 Objectives", level=2)
body("The principal objectives of the project are as follows:")
bullets([
    "To design and implement a role-based access control model that clearly separates the "
    "privileges of Admin, Teacher, and Student users.",
    "To enable teachers to create, update, and publish courses composed of structured video lessons.",
    "To allow students to browse a searchable course catalogue, enrol in courses, and track "
    "their learning progress.",
    "To provide administrators with tools for managing user accounts and moderating published courses.",
    "To deliver a secure and responsive system that satisfies the defined non-functional "
    "requirements for performance and data protection.",
    "To apply a disciplined software engineering process, including requirements engineering, "
    "version control, and iterative delivery, within a nine-week schedule.",
])

doc.add_heading("1.4 Academic and Technical References", level=2)
refs = [
    "Sommerville, I. (2016). Software Engineering (10th ed.). Pearson.",
    "Pressman, R. S., & Maxim, B. R. (2020). Software Engineering: A Practitioner's Approach "
    "(9th ed.). McGraw-Hill Education.",
    "ISO/IEC/IEEE 29148:2018. Systems and Software Engineering – Life Cycle Processes – "
    "Requirements Engineering. International Organization for Standardization.",
    "Sandhu, R. S., Coyne, E. J., Feinstein, H. L., & Youman, C. E. (1996). Role-based access "
    "control models. IEEE Computer, 29(2), 38–47.",
    "Wiegers, K., & Beatty, J. (2013). Software Requirements (3rd ed.). Microsoft Press.",
    "OWASP Foundation. Password Storage Cheat Sheet. OWASP Cheat Sheet Series. "
    "https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html",
    "Google. Web Vitals. web.dev. https://web.dev/articles/vitals",
]
for i, ref in enumerate(refs, 1):
    p = doc.add_paragraph(f"[{i}] {ref}")
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- Week 2 ----------
doc.add_heading("Part II: Week 2 – Requirements and Planning", level=1)

doc.add_heading("2.1 Functional Requirements", level=2)
table(["ID", "Requirement"], [
    ("FR-01", "The system shall allow users to register and log in as Admin, Teacher, or Student."),
    ("FR-02", "The system shall allow administrators to manage platform users and moderate courses."),
    ("FR-03", "The system shall allow teachers to create, update, and publish course content and video lessons."),
    ("FR-04", "The system shall allow students to browse the course catalog and enroll in available courses."),
    ("FR-05", "The system shall allow enrolled students to watch video lessons and track their progress."),
], [2.2, 13.8])

doc.add_heading("2.2 Non-Functional Requirements", level=2)
table(["ID", "Category", "Requirement"], [
    ("NFR-01", "Performance",
     "The system shall load course catalog pages in under 2 seconds under normal network conditions."),
    ("NFR-02", "Security",
     "User passwords must be securely hashed before being stored in the database."),
], [2.2, 3.0, 10.8])

doc.add_heading("2.3 Project Scope", level=2)
para("In Scope", bold=True, after=2)
bullets([
    "Management of the three user roles (Admin, Teacher, Student) and their respective permissions.",
    "Course creation, editing, and publishing by teachers.",
    "Student registration, course enrolment, and viewing of course content.",
    "Tracking of lesson completion and overall course progress for enrolled students.",
])
para("Out of Scope", bold=True, after=2)
bullets([
    "Online payment gateways and paid course transactions.",
    "Live (real-time) video streaming.",
])

doc.add_heading("2.4 Constraints", level=2)
bullets([
    "Time: the project must be completed and submitted within the academic semester (nine weeks).",
    "Resources: development is carried out by a small student team using free or open-source "
    "tools and services.",
    "Hosting: video content is limited to pre-recorded files within the storage limits of the "
    "selected hosting environment.",
])

doc.add_heading("2.5 Assumptions", level=2)
bullets([
    "Users have access to a modern web browser and a stable internet connection.",
    "Teachers provide their own pre-recorded video lessons in commonly supported formats.",
    "Each user account is associated with exactly one role at any given time.",
    "Administrators are trusted users responsible for approving or removing inappropriate content.",
])

doc.add_heading("2.6 Nine-Week Project Milestones", level=2)
table(["Week", "Milestone", "Main Deliverables", "Status"], [
    ("1", "Project definition", "Title, problem statement, objectives, references", "Completed"),
    ("2", "Requirements and planning", "FR/NFR, scope, constraints, assumptions, milestones", "Completed"),
    ("3", "System design", "Use case diagram, ER diagram, system architecture", "Planned"),
    ("4", "UI/UX design", "Wireframes and page flow for each role", "Planned"),
    ("5", "Authentication and roles", "Registration, login, role-based access (FR-01)", "Planned"),
    ("6", "Course management", "Teacher course creation and admin moderation (FR-02, FR-03)", "Planned"),
    ("7", "Learning features", "Catalogue, enrolment, video lessons, progress (FR-04, FR-05)", "Planned"),
    ("8", "Testing and refinement", "Functional, performance and security testing (NFR-01, NFR-02)", "Planned"),
    ("9", "Final delivery", "Deployment, final report and presentation", "Planned"),
], [1.5, 3.8, 8.0, 2.4])

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
doc.add_heading("Part III: Project Repository", level=1)
body("The source code, documentation, and task tracking for the project are maintained on GitHub:")
table(["Item", "Link"], [
    ("Repository URL", REPO),
    ("Issues Page", f"{REPO}/issues"),
    ("Commit History", f"{REPO}/commits/main"),
], [4.0, 12.0])

footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("SkillStream – Graduation Project Report (Week 1 & 2) – Altınbaş University")
fr.font.size = Pt(9)
fr.italic = True

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(f"Saved {OUT}")
